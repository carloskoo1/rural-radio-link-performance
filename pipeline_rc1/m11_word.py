
from __future__ import annotations

from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .m00_config import TABLE_TITLES, FIGURE_TITLES, FIGURE_START


def _border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "B7B7B7")
        borders.append(e)
    tc_pr.append(borders)


def _shade(cell) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9EAF7")
    cell._tc.get_or_add_tcPr().append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)



def _set_keep_with_next(paragraph_obj, value: bool = True) -> None:
    paragraph_obj.paragraph_format.keep_with_next = value


def _set_keep_together(paragraph_obj, value: bool = True) -> None:
    paragraph_obj.paragraph_format.keep_together = value


def _landscape_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)
    return section


def _portrait_section(doc: Document):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    return section

def paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10)


def add_table(
    doc: Document,
    key: str,
    df: pd.DataFrame,
    introduction: str,
    interpretation: str,
    note: str | None = None,
) -> None:
    paragraph(doc, introduction)

    p = doc.add_paragraph()
    _set_keep_with_next(p)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(TABLE_TITLES[key])
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _repeat_header(table.rows[0])

    for i, column in enumerate(df.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(column)
        _shade(cell)
        _border(cell)
        for pp in cell.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for rr in pp.runs:
                rr.bold = True
                rr.font.size = Pt(8)

    for _, row in df.iterrows():
        word_row = table.add_row()
        _prevent_split(word_row)
        for i, value in enumerate(row):
            cell = word_row.cells[i]
            if pd.isna(value):
                text = ""
            elif isinstance(value, float):
                text = f"{value:.3f}".rstrip("0").rstrip(".")
            else:
                text = str(value)
            cell.text = text
            _border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for pp in cell.paragraphs:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs:
                    rr.font.size = Pt(7.5)

    if note:
        p = doc.add_paragraph()
        r = p.add_run(note)
        r.italic = True
        r.font.size = Pt(8)

    paragraph(doc, interpretation)


def add_figure(
    doc: Document,
    path: Path,
    number: int,
    title: str,
    introduction: str,
    interpretation: str,
) -> None:
    # Cada figura inicia en una zona limpia para evitar particiones.
    if len(doc.paragraphs) > 0:
        doc.add_page_break()

    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.space_after = Pt(6)
    _set_keep_with_next(p_intro)
    r = p_intro.add_run(introduction)
    r.font.name = "Arial"
    r.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    _set_keep_with_next(p)
    p.add_run().add_picture(str(path), width=Cm(16.2))

    p_caption = doc.add_paragraph()
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_caption.paragraph_format.space_after = Pt(2)
    _set_keep_with_next(p_caption)
    r = p_caption.add_run(f"Fig. {number}. {title}")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)

    p_source = doc.add_paragraph()
    p_source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_source.paragraph_format.space_after = Pt(6)
    _set_keep_with_next(p_source)
    r = p_source.add_run("Fuente: Elaboración propia.")
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(8)

    p_inter = doc.add_paragraph()
    p_inter.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_inter.paragraph_format.space_after = Pt(8)
    p_inter.paragraph_format.line_spacing = 1.08
    _set_keep_together(p_inter)
    r = p_inter.add_run(interpretation)
    r.font.name = "Arial"
    r.font.size = Pt(10)


def create_word(
    output_path: Path,
    tables: dict[str, pd.DataFrame],
    texts: dict[str, tuple[str, str]],
    figures: dict[str, Path],
) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CAPÍTULO IV\nRESULTADOS")
    run.bold = True
    run.font.size = Pt(14)

    # 4.1 Normalidad
    doc.add_heading("4.1. Verificación de normalidad y selección de pruebas", level=1)
    add_table(
        doc, "normalidad", tables["normalidad"], *texts["normalidad"],
        note=(
            "Nota. La selección no paramétrica se sustentó en la prueba formal, "
            "la naturaleza discreta del MCS y la inspección de las distribuciones."
        ),
    )
    add_figure(
        doc, figures["normalidad"], 7,
        "Estadísticos W de Shapiro-Wilk por métrica y escenario.",
        "La Figura 7 complementa la Tabla IV y permite comparar la intensidad de "
        "las desviaciones respecto de la normalidad.",
        "Los valores W se mantuvieron por debajo de uno en todas las combinaciones. "
        "Las desviaciones fueron particularmente marcadas para el MCS de E0 y para "
        "algunas distribuciones de RSSI y SNR."
    )

    # 4.2 Inferencia
    doc.add_heading("4.2. Contraste inferencial entre escenarios", level=1)
    add_table(doc, "kruskal", tables["kruskal"], *texts["kruskal"])
    add_figure(
        doc, figures["kruskal"], 8,
        "Tamaños del efecto de las diferencias entre escenarios.",
        "La Figura 8 representa los valores de ε² obtenidos mediante Kruskal-Wallis.",
        "RSSI presentó el mayor tamaño del efecto, seguido de SNR, MCS y throughput. "
        "En todos los casos la magnitud fue grande."
    )
    add_table(doc, "dunn", tables["dunn"], *texts["dunn"])
    add_figure(
        doc, figures["dunn"], 9,
        "Proporción de comparaciones post hoc significativas.",
        "La Figura 9 sintetiza la proporción de pares diferenciados mediante Dunn.",
        "La alta proporción de comparaciones significativas confirma que las "
        "diferencias globales no estuvieron concentradas en un único par de escenarios."
    )

    # 4.3 Descriptivos
    doc.add_heading("4.3. Caracterización descriptiva", level=1)
    _landscape_section(doc)
    add_table(doc, "descriptivos", tables["descriptivos"], *texts["descriptivos"])
    _portrait_section(doc)

    add_figure(
        doc, figures["rssi_media"], 10,
        "RSSI promedio y desviación estándar por escenario.",
        "La Figura 10 representa conjuntamente el nivel medio y la dispersión del RSSI.",
        "E2 presentó la señal recibida menos negativa, mientras que E4 registró la "
        "menor dispersión. Esto evidencia que intensidad y estabilidad no coincidieron "
        "necesariamente en un mismo escenario."
    )
    add_figure(
        doc, figures["snr_media"], 11,
        "SNR promedio y desviación estándar por escenario.",
        "La Figura 11 muestra el comportamiento medio del SNR y su dispersión.",
        "E2 alcanzó el mayor SNR medio, mientras que E3 presentó la menor dispersión "
        "del indicador."
    )
    add_figure(
        doc, figures["mcs_media"], 12,
        "MCS promedio y desviación estándar por escenario.",
        "La Figura 12 permite comparar el nivel medio de modulación y su dispersión.",
        "Los promedios de MCS fueron próximos, pero su variabilidad diferenció con "
        "claridad a los escenarios."
    )

    add_table(
        doc, "throughput", tables["throughput"], *texts["throughput"],
        note=(
            "Nota. El throughput corresponde al tráfico observado durante el "
            "monitoreo y no a la capacidad nominal del enlace."
        ),
    )
    add_figure(
        doc, figures["throughput"], 13,
        "Throughput descendente medio y desviación estándar.",
        "La Figura 13 representa el tráfico observado y su variabilidad.",
        "E4 registró el mayor promedio, aunque E3 presentó la mayor mediana y una "
        "variabilidad relativa inferior."
    )

    # 4.4 Franjas
    doc.add_heading("4.4. Comparación por franjas horarias", level=1)
    add_table(
        doc, "franjas", tables["franjas"], *texts["franjas"],
        note=(
            "Nota. Los contrastes se realizaron dentro de cada escenario; los "
            "valores p se ajustaron mediante Holm."
        ),
    )
    add_figure(
        doc, figures["franjas"], 14,
        "Escenarios con diferencias significativas entre franjas horarias.",
        "La Figura 14 resume el número de escenarios diferenciados para cada métrica.",
        "MCS y throughput presentaron diferencias en todos los escenarios con datos "
        "válidos, mientras que RSSI y SNR mostraron un patrón menos generalizado."
    )
    add_figure(
        doc, figures["cliff"], 15,
        "Delta de Cliff por métrica y escenario.",
        "La Figura 15 representa la dirección y magnitud de los efectos temporales.",
        "Los efectos más grandes se concentraron en el throughput. Los valores "
        "negativos del MCS indican niveles superiores durante la franja 01:00–03:00 "
        "respecto de 20:00–22:00."
    )

    # 4.5 Precipitación
    doc.add_heading("4.5. Asociación con la precipitación", level=1)
    add_table(
        doc, "precipitacion_diag", tables["precipitacion_diag"],
        "Antes de interpretar la precipitación, se examinó la consistencia de la "
        "variable consignada como precip_mm.",
        "El diagnóstico permitió verificar la distribución y la consistencia horaria "
        "de la covariable antes de los análisis de asociación.",
        note="Nota. La unidad, la columna de entrada y el factor de conversión aplicado se presentan explícitamente en la tabla.",
    )
    add_table(
        doc, "precipitacion_sens", tables["precipitacion_sens"],
        "Se evaluó la sensibilidad de la clasificación frente a distintos umbrales.",
        "La proporción de registros clasificados con lluvia varió de forma relevante "
        "según el umbral; por ello, el criterio de 0.10 mm/h se mantuvo explícito.",
    )
    add_figure(
        doc, figures["lluvia_sens"], 16,
        "Sensibilidad de la clasificación de precipitación.",
        "La Figura 16 muestra la variación de la proporción de lluvia al modificar el umbral.",
        "El porcentaje disminuyó progresivamente al elevar el umbral. El valor de "
        "0.10 mm/h permitió mantener una definición explícita y reproducible."
    )
    add_table(
        doc, "precipitacion_conteo", tables["precipitacion_conteo"],
        "Con el umbral seleccionado, se cuantificaron los registros con lluvia y sin lluvia.",
        "La distribución resultante permitió comparar ambos estratos definidos por el "
        "umbral, aunque con predominio de registros iguales o superiores a 0.10 mm/h.",
    )
    add_figure(
        doc, figures["lluvia_conteo"], 17,
        "Distribución de registros según condición de precipitación.",
        "La Figura 17 representa la composición porcentual de los registros climáticos.",
        "El estrato con precipitación igual o superior a 0.10 mm/h concentró la mayor "
        "parte de las observaciones, por lo que las comparaciones deben "
        "interpretarse considerando este desbalance."
    )
    add_table(
        doc, "precipitacion_senal", tables["precipitacion_senal"],
        "Se compararon RSSI y SNR entre los estratos definidos por el umbral de 0.10 mm/h.",
        "Las diferencias descriptivas fueron reducidas y se interpretaron junto con "
        "el análisis correlacional horario.",
    )
    add_table(
        doc, "precipitacion_capacidad", tables["precipitacion_capacidad"],
        "Se examinó el comportamiento de MCS y throughput en ambos estratos de precipitación.",
        "La proximidad de los estadísticos descriptivos no evidenció una separación "
        "amplia entre las condiciones atmosféricas.",
    )
    add_figure(
        doc, figures["lluvia_throughput"], 18,
        "Throughput observado según condición de precipitación.",
        "La Figura 18 compara el throughput medio bajo ambas condiciones.",
        "Los promedios fueron próximos y los intervalos de dispersión se superpusieron, "
        "lo que es consistente con una asociación climática de magnitud reducida."
    )
    add_table(doc, "spearman", tables["spearman"], *texts["precipitacion"])
    add_figure(
        doc, figures["spearman"], 19,
        "Correlaciones de Spearman entre precipitación y métricas.",
        "La Figura 19 sintetiza la dirección y magnitud de las asociaciones.",
        "Todos los coeficientes se mantuvieron en el rango de asociaciones débiles. "
        "Throughput presentó la mayor asociación positiva y RSSI la mayor asociación negativa."
    )

    # 4.6 Estabilidad
    doc.add_heading("4.6. Estabilidad operativa", level=1)
    add_table(
        doc, "estabilidad_senal", tables["estabilidad_senal"],
        "La estabilidad de RSSI se evaluó mediante DE y RIC; para SNR se empleó CV.",
        "La comparación evidenció que la configuración de mayor estabilidad dependió "
        "del indicador considerado.",
    )
    add_figure(
        doc, figures["est_rssi"], 20,
        "Desviación estándar del RSSI descendente.",
        "La Figura 20 complementa la Tabla XV-A mediante la comparación de la dispersión.",
        "E4 presentó la menor desviación estándar del RSSI, mientras que E0 registró "
        "la mayor dispersión."
    )
    add_table(
        doc, "estabilidad_capacidad", tables["estabilidad_capacidad"],
        "La estabilidad de MCS y throughput se evaluó mediante coeficientes de variación.",
        "Los resultados confirmaron que ningún escenario optimizó simultáneamente "
        "todas las métricas de estabilidad.",
    )
    add_figure(
        doc, figures["est_mcs"], 21,
        "Coeficiente de variación del MCS descendente.",
        "La Figura 21 muestra la variabilidad relativa de la modulación.",
        "E5 presentó el menor CV del MCS, seguido de E3."
    )
    add_figure(
        doc, figures["est_throughput"], 22,
        "Coeficiente de variación del throughput descendente.",
        "La Figura 22 muestra la variabilidad relativa del tráfico observado.",
        "E3 registró el menor CV del throughput, mientras que E2 presentó el mayor."
    )
    add_figure(
        doc, figures["indice"], 23,
        "Índice compuesto de estabilidad operativa.",
        "Los indicadores individuales se integraron en un índice compuesto de estabilidad.",
        texts["estabilidad"][1],
    )
    paragraph(
        doc,
        "El índice compuesto se calculó como I_est = 1 - promedio(V_j*), donde V_j* "
        "representa la normalización Min-Max de la DE del RSSI y de los CV de SNR, "
        "MCS y throughput. Sus valores constituyen una comparación interna del "
        "conjunto de escenarios y no una escala universal."
    )

    # 4.7 Síntesis
    doc.add_heading("4.7. Síntesis integrada", level=1)
    add_table(
        doc, "objetivos", tables["objetivos"],
        "Los hallazgos fueron organizados según los objetivos específicos y la hipótesis general.",
        "La integración de la evidencia permitió atender los cuatro objetivos y "
        "respaldar la hipótesis general dentro de las condiciones evaluadas.",
    )
    paragraph(
        doc,
        "En términos integrales, las configuraciones experimentales se asociaron "
        "con diferencias estadísticamente significativas y tamaños del efecto grandes "
        "en RSSI, SNR, MCS y throughput. El análisis por franjas horarias evidenció "
        "diferencias especialmente consistentes en MCS y throughput, mientras que "
        "la precipitación presentó asociaciones débiles con las métricas técnicas. "
        "La evaluación de estabilidad mostró que ningún escenario optimizó "
        "simultáneamente todos los indicadores; sin embargo, el índice compuesto "
        "identificó a E3 como la configuración de mayor estabilidad relativa. En "
        "consecuencia, los resultados respaldan la hipótesis general dentro de los "
        "periodos y condiciones evaluados, sin atribuir causalidad absoluta."
    )

    # Anexos
    annex = doc.add_section(WD_SECTION.NEW_PAGE)
    annex.orientation = WD_ORIENT.LANDSCAPE
    annex.page_width, annex.page_height = annex.page_height, annex.page_width
    annex.top_margin = Cm(1.5)
    annex.bottom_margin = Cm(1.5)
    annex.left_margin = Cm(1.3)
    annex.right_margin = Cm(1.3)

    doc.add_heading("ANEXOS ESTADÍSTICOS", level=1)
    add_table(
        doc, "anexo_dunn", tables["anexo_dunn"],
        "La Tabla A1 presenta la totalidad de las comparaciones de Dunn.",
        "El anexo constituye el respaldo íntegro de la síntesis presentada en la Tabla VI.",
    )
    add_table(
        doc, "anexo_franjas", tables["anexo_franjas"],
        "La Tabla A2 presenta el detalle de las comparaciones por franja y escenario.",
        "El anexo permite auditar los valores p ajustados y los tamaños del efecto.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
