"""
11_generar_capitulo_iv.py
=========================

Genera automáticamente:
- outputs/reportes/Capitulo_IV_Resultados.docx
- outputs/reportes/resultados_tesis_v2.xlsx

Mejoras respecto de la versión anterior:
- limpia encabezados multinivel;
- redondea valores;
- usa títulos legibles;
- evita figuras duplicadas;
- presenta Dunn resumido en el cuerpo;
- envía Dunn completo a un anexo;
- genera interpretación automática;
- utiliza numeración coherente de tablas y figuras.

Requisitos:
    python -m pip install pandas xlsxwriter python-docx pillow

Ejecución:
    python scripts/11_generar_capitulo_iv.py
"""

from __future__ import annotations

import hashlib
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BASE_DIR = Path(__file__).resolve().parent.parent
TABLAS_DIR = BASE_DIR / "outputs" / "tables"
FIGURAS_DIR = BASE_DIR / "outputs" / "figures"
REPORTES_DIR = BASE_DIR / "outputs" / "reportes"

DOCX_SALIDA = REPORTES_DIR / "Capitulo_IV_Resultados.docx"
XLSX_SALIDA = REPORTES_DIR / "resultados_tesis_v2.xlsx"

ALPHA = 0.05


TITULOS_COLUMNAS = {
    "escenario": "Escenario",
    "frecuencia_mhz": "Frecuencia (MHz)",
    "ancho_canal_mhz": "Canal (MHz)",
    "n_registros": "N",
    "dias_unicos": "Días",
    "rssi_dl_media": "RSSI media (dBm)",
    "rssi_dl_de": "RSSI DE",
    "snr_dl_media": "SNR media (dB)",
    "snr_dl_de": "SNR DE",
    "mcs_dl_media": "MCS media",
    "mcs_dl_de": "MCS DE",
    "condicion_lluvia": "Condición",
    "porcentaje": "Porcentaje (%)",
    "variable": "Métrica",
    "rho_spearman": "ρ de Spearman",
    "n": "N",
    "rssi_dl_mean": "RSSI media (dBm)",
    "rssi_dl_std": "RSSI DE",
    "rssi_dl_cv": "RSSI CV (%)",
    "snr_dl_mean": "SNR media (dB)",
    "snr_dl_std": "SNR DE",
    "snr_dl_cv": "SNR CV (%)",
    "mcs_dl_mean": "MCS media",
    "mcs_dl_std": "MCS DE",
    "mcs_dl_cv": "MCS CV (%)",
    "throughput_dl_mean": "Throughput media (Mbps)",
    "throughput_dl_std": "Throughput DE",
    "throughput_dl_cv": "Throughput CV (%)",
}

NOMBRES_METRICAS = {
    "rssi_dl": "RSSI DL",
    "snr_dl": "SNR DL",
    "mcs_dl": "MCS DL",
    "throughput_dl": "Throughput DL",
}

FIGURAS = {
    "fig_4_9_cv_rssi": (
        "Fig. 7. Coeficiente de variación del RSSI descendente por escenario experimental."
    ),
    "fig_4_10_boxplot_rssi": (
        "Fig. 8. Distribución del RSSI descendente por escenario experimental."
    ),
    "fig_4_11_cv_mcs": (
        "Fig. 9. Coeficiente de variación del MCS descendente por escenario experimental."
    ),
    "fig_4_11_cv_throughput": (
        "Fig. 10. Coeficiente de variación del throughput descendente observado."
    ),
    "fig_4_12_indice_variabilidad": (
        "Fig. 11. Índice relativo de variabilidad operativa por escenario experimental."
    ),
}


def leer_csv_normal(ruta: Path) -> pd.DataFrame:
    return pd.read_csv(ruta, encoding="utf-8-sig")


def leer_tabla_lluvia_multinivel(ruta: Path) -> pd.DataFrame:
    """Convierte el CSV multinivel de desempeño según lluvia a columnas planas."""
    bruto = pd.read_csv(ruta, header=[0, 1], index_col=0, encoding="utf-8-sig")
    bruto.index.name = "Condición"

    columnas = []
    for metrica, estadistico in bruto.columns:
        metrica_legible = {
            "rssi_dl": "RSSI",
            "snr_dl": "SNR",
            "mcs_dl": "MCS",
            "throughput_dl": "Throughput",
        }.get(str(metrica), str(metrica))
        estadistico_legible = {
            "mean": "media",
            "std": "DE",
            "median": "mediana",
        }.get(str(estadistico), str(estadistico))
        columnas.append(f"{metrica_legible} {estadistico_legible}")

    bruto.columns = columnas
    return bruto.reset_index()


def limpiar_tabla(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df = df.dropna(how="all")
    df = df.rename(columns=TITULOS_COLUMNAS)

    for columna in df.columns:
        if pd.api.types.is_numeric_dtype(df[columna]):
            df[columna] = df[columna].round(3)

    if "Métrica" in df.columns:
        df["Métrica"] = df["Métrica"].replace(NOMBRES_METRICAS)

    return df


def cargar_tablas() -> dict[str, pd.DataFrame]:
    rutas = {
        "descriptivos": TABLAS_DIR / "tabla_4_1_descriptivos_principales.csv",
        "kruskal": TABLAS_DIR / "tabla_4_2_kruskal_wallis.csv",
        "lluvia_conteo": TABLAS_DIR / "tabla_4_3_conteo_lluvia.csv",
        "dunn_completa": TABLAS_DIR / "tabla_4_3_dunn_completa.csv",
        "dunn_significativas": TABLAS_DIR / "tabla_4_3_dunn_significativas.csv",
        "lluvia_desempeno": TABLAS_DIR / "tabla_4_4_desempeno_lluvia.csv",
        "spearman": TABLAS_DIR / "tabla_4_5_correlacion_precipitacion.csv",
        "estabilidad": TABLAS_DIR / "tabla_4_6_estabilidad_por_escenario.csv",
    }

    faltantes = [str(ruta) for ruta in rutas.values() if not ruta.exists()]
    if faltantes:
        raise FileNotFoundError(
            "Faltan archivos requeridos:\n- " + "\n- ".join(faltantes)
        )

    tablas = {
        "descriptivos": limpiar_tabla(leer_csv_normal(rutas["descriptivos"])),
        "kruskal": limpiar_tabla(leer_csv_normal(rutas["kruskal"])),
        "lluvia_conteo": limpiar_tabla(leer_csv_normal(rutas["lluvia_conteo"])),
        "dunn_completa": limpiar_tabla(leer_csv_normal(rutas["dunn_completa"])),
        "dunn_significativas": limpiar_tabla(
            leer_csv_normal(rutas["dunn_significativas"])
        ),
        "lluvia_desempeno": limpiar_tabla(
            leer_tabla_lluvia_multinivel(rutas["lluvia_desempeno"])
        ),
        "spearman": limpiar_tabla(leer_csv_normal(rutas["spearman"])),
        "estabilidad": limpiar_tabla(leer_csv_normal(rutas["estabilidad"])),
    }

    return tablas


def resumen_dunn(tabla_completa: pd.DataFrame) -> pd.DataFrame:
    filas = []

    for metrica, grupo in tabla_completa.groupby("Métrica", sort=False):
        total = len(grupo)
        significativas = int(
            (grupo["Diferencia significativa"].astype(str).str.strip() == "Sí").sum()
        )
        no_sig = grupo.loc[
            grupo["Diferencia significativa"].astype(str).str.strip() != "Sí"
        ]

        if no_sig.empty:
            pares = "Ninguna"
        else:
            pares = "; ".join(
                f"{fila['Comparación']} (p={fila['Valor p ajustado']})"
                for _, fila in no_sig.iterrows()
            )

        filas.append(
            {
                "Métrica": metrica,
                "Comparaciones significativas": f"{significativas} de {total}",
                "Comparaciones no significativas": pares,
            }
        )

    return pd.DataFrame(filas)


def hash_archivo(ruta: Path) -> str:
    digest = hashlib.sha256()
    with ruta.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def listar_figuras_unicas() -> list[tuple[Path, str]]:
    if not FIGURAS_DIR.exists():
        return []

    candidatas = sorted(
        [
            ruta
            for ruta in FIGURAS_DIR.iterdir()
            if ruta.suffix.lower() in {".png", ".jpg", ".jpeg"}
        ]
    )

    vistas: set[str] = set()
    seleccionadas: list[tuple[Path, str]] = []

    for ruta in candidatas:
        clave = hash_archivo(ruta)
        if clave in vistas:
            continue
        vistas.add(clave)

        titulo = FIGURAS.get(ruta.stem)
        if titulo:
            seleccionadas.append((ruta, titulo))

    return seleccionadas


def bordes(celda) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elemento = OxmlElement(f"w:{lado}")
        elemento.set(qn("w:val"), "single")
        elemento.set(qn("w:sz"), "4")
        elemento.set(qn("w:color"), "B7B7B7")
        tc_borders.append(elemento)
    tc_pr.append(tc_borders)


def sombrear(celda, color: str = "D9EAF7") -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def agregar_parrafo(documento: Document, texto: str, negrita: bool = False) -> None:
    parrafo = documento.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = parrafo.add_run(texto)
    run.bold = negrita
    run.font.name = "Arial"
    run.font.size = Pt(10)


def agregar_tabla(
    documento: Document,
    titulo: str,
    df: pd.DataFrame,
    nota: str | None = None,
    max_filas: int | None = None,
) -> None:
    parrafo = documento.add_paragraph()
    run = parrafo.add_run(titulo)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)

    vista = df.head(max_filas) if max_filas else df

    tabla = documento.add_table(rows=1, cols=len(vista.columns))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = True

    for i, columna in enumerate(vista.columns):
        celda = tabla.rows[0].cells[i]
        celda.text = str(columna)
        sombrear(celda)
        bordes(celda)
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in celda.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(8)

    for _, fila in vista.iterrows():
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = "" if pd.isna(valor) else str(valor)
            bordes(celdas[i])
            celdas[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in celdas[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(8)

    if nota:
        p = documento.add_paragraph()
        run = p.add_run(nota)
        run.italic = True
        run.font.name = "Arial"
        run.font.size = Pt(8)

    documento.add_paragraph()


def interpretar_descriptivos(df: pd.DataFrame) -> str:
    rssi_col = "RSSI media (dBm)"
    snr_col = "SNR media (dB)"
    mcs_col = "MCS media"

    mejor_rssi = df.loc[df[rssi_col].idxmax()]
    mejor_snr = df.loc[df[snr_col].idxmax()]
    mejor_mcs = df.loc[df[mcs_col].idxmax()]

    return (
        f"Descriptivamente, el escenario {mejor_rssi['Escenario']} presentó el "
        f"mejor RSSI promedio ({mejor_rssi[rssi_col]:.2f} dBm), mientras que "
        f"{mejor_snr['Escenario']} alcanzó el mayor SNR "
        f"({mejor_snr[snr_col]:.2f} dB). El mayor MCS promedio correspondió a "
        f"{mejor_mcs['Escenario']} ({mejor_mcs[mcs_col]:.2f}). Estas diferencias "
        "fueron posteriormente contrastadas mediante procedimientos no paramétricos."
    )


def interpretar_kruskal(df: pd.DataFrame) -> str:
    metricas = ", ".join(df["Métrica"].astype(str).tolist())
    return (
        f"La prueba de Kruskal-Wallis identificó diferencias estadísticamente "
        f"significativas para {metricas}, con valores p inferiores a 0.001. "
        "Los tamaños del efecto fueron clasificados como grandes, lo que indicó "
        "una influencia relevante de la configuración técnica sobre las "
        "distribuciones de las métricas evaluadas."
    )


def interpretar_lluvia(conteo: pd.DataFrame, desempeno: pd.DataFrame) -> str:
    lluvia = conteo.loc[conteo["Condición"] == "Con lluvia"].iloc[0]
    return (
        f"El {lluvia['Porcentaje (%)']:.2f} % de los registros se obtuvo bajo "
        "condiciones de lluvia. Debido al fuerte desequilibrio entre las categorías "
        "con lluvia y sin lluvia, las diferencias descriptivas se interpretaron "
        "con cautela y el análisis principal de asociación se realizó mediante "
        "el coeficiente de Spearman."
    )


def interpretar_spearman(df: pd.DataFrame) -> str:
    max_abs = df.loc[df["ρ de Spearman"].abs().idxmax()]
    return (
        f"Las asociaciones entre precipitación y desempeño fueron débiles. La "
        f"mayor magnitud correspondió a {max_abs['Métrica']} "
        f"(ρ={max_abs['ρ de Spearman']:.3f}), por lo que la precipitación no se "
        "consideró un factor determinante del comportamiento técnico observado."
    )


def interpretar_estabilidad(df: pd.DataFrame) -> str:
    mejor_rssi = df.loc[df["RSSI CV (%)"].idxmin()]
    mejor_mcs = df.loc[df["MCS CV (%)"].idxmin()]
    throughput = df.dropna(subset=["Throughput CV (%)"])
    mejor_th = throughput.loc[throughput["Throughput CV (%)"].idxmin()]

    return (
        f"La menor variabilidad relativa del RSSI se observó en "
        f"{mejor_rssi['Escenario']} ({mejor_rssi['RSSI CV (%)']:.3f} %); "
        f"la menor variabilidad del MCS correspondió a "
        f"{mejor_mcs['Escenario']} ({mejor_mcs['MCS CV (%)']:.3f} %); y el "
        f"throughput más regular se registró en {mejor_th['Escenario']} "
        f"({mejor_th['Throughput CV (%)']:.3f} %). Por tanto, la estabilidad "
        "operativa dependió de la métrica evaluada y no de un único escenario."
    )


def agregar_figuras(documento: Document, figuras: list[tuple[Path, str]]) -> None:
    for ruta, titulo in figuras:
        p = documento.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ruta), width=Cm(15.5))

        p2 = documento.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(titulo)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)

        p3 = documento.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p3.add_run("Fuente: Elaboración propia.")
        r.italic = True
        r.font.name = "Arial"
        r.font.size = Pt(8)

        if "variabilidad operativa" in titulo.lower():
            agregar_parrafo(
                documento,
                "Valores menores del índice representan menor variabilidad "
                "operativa y, por tanto, una mayor estabilidad relativa.",
            )


def crear_word(tablas: dict[str, pd.DataFrame]) -> None:
    documento = Document()
    seccion = documento.sections[0]
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(2.5)
    seccion.right_margin = Cm(2.5)

    documento.styles["Normal"].font.name = "Arial"
    documento.styles["Normal"].font.size = Pt(10)

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("CAPÍTULO IV\nRESULTADOS")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14)

    documento.add_heading("4.1. Verificación de supuestos y análisis inferencial", level=1)
    agregar_parrafo(
        documento,
        "La selección de las pruebas inferenciales se sustentó en la evaluación "
        "previa de normalidad. Debido al incumplimiento de este supuesto, se "
        "emplearon procedimientos no paramétricos para comparar los escenarios."
    )

    agregar_tabla(
        documento,
        "Tabla II. Resultados de la prueba de Kruskal-Wallis para las métricas de desempeño técnico",
        tablas["kruskal"],
        nota=(
            "Nota. El escenario E0 no se incluyó en el contraste de throughput "
            "por no disponer de registros equivalentes a cnMaestro."
        ),
    )
    agregar_parrafo(documento, interpretar_kruskal(tablas["kruskal"]))

    dunn_resumen = resumen_dunn(tablas["dunn_completa"])
    agregar_tabla(
        documento,
        "Tabla III. Síntesis de las comparaciones post hoc de Dunn",
        dunn_resumen,
        nota=(
            "Nota. Los valores p fueron ajustados mediante el procedimiento de Holm."
        ),
    )
    agregar_parrafo(
        documento,
        "Las comparaciones post hoc mostraron que la influencia de la configuración "
        "no implicó diferencias entre todos los pares. Se identificaron equivalencias "
        "estadísticas puntuales, cuya interpretación se realizó conjuntamente con "
        "las medianas y la relevancia técnica de las diferencias."
    )

    documento.add_heading("4.2. Comparación descriptiva entre escenarios", level=1)
    agregar_tabla(
        documento,
        "Tabla IV. Estadísticos descriptivos principales por escenario experimental",
        tablas["descriptivos"],
    )
    agregar_parrafo(documento, interpretar_descriptivos(tablas["descriptivos"]))

    documento.add_heading("4.3. Precipitación y desempeño técnico", level=1)
    agregar_tabla(
        documento,
        "Tabla V. Distribución de registros según condición de precipitación",
        tablas["lluvia_conteo"],
    )
    agregar_tabla(
        documento,
        "Tabla VI. Desempeño técnico según condición de precipitación",
        tablas["lluvia_desempeno"],
    )
    agregar_parrafo(
        documento,
        interpretar_lluvia(tablas["lluvia_conteo"], tablas["lluvia_desempeno"]),
    )

    agregar_tabla(
        documento,
        "Tabla VII. Correlación entre precipitación y métricas técnicas",
        tablas["spearman"],
    )
    agregar_parrafo(documento, interpretar_spearman(tablas["spearman"]))

    documento.add_heading("4.4. Estabilidad operativa", level=1)
    agregar_tabla(
        documento,
        "Tabla VIII. Indicadores de estabilidad operativa por escenario",
        tablas["estabilidad"],
    )
    agregar_parrafo(documento, interpretar_estabilidad(tablas["estabilidad"]))

    figuras = listar_figuras_unicas()
    agregar_figuras(documento, figuras)

    documento.add_page_break()
    documento.add_heading(
        "ANEXO. Comparaciones post hoc de Dunn completas",
        level=1,
    )

    seccion_anexo = documento.add_section(start_type=1)
    seccion_anexo.orientation = WD_ORIENT.LANDSCAPE
    seccion_anexo.page_width, seccion_anexo.page_height = (
        seccion_anexo.page_height,
        seccion_anexo.page_width,
    )
    seccion_anexo.top_margin = Cm(1.5)
    seccion_anexo.bottom_margin = Cm(1.5)
    seccion_anexo.left_margin = Cm(1.3)
    seccion_anexo.right_margin = Cm(1.3)

    agregar_tabla(
        documento,
        "Tabla A. Comparaciones post hoc de Dunn completas",
        tablas["dunn_completa"],
        nota="Nota. Los valores p fueron ajustados mediante el procedimiento de Holm.",
    )

    documento.save(DOCX_SALIDA)


def crear_excel(tablas: dict[str, pd.DataFrame]) -> None:
    hojas = {
        "Descriptivos": tablas["descriptivos"],
        "Kruskal-Wallis": tablas["kruskal"],
        "Dunn resumen": resumen_dunn(tablas["dunn_completa"]),
        "Dunn completa": tablas["dunn_completa"],
        "Dunn significativas": tablas["dunn_significativas"],
        "Conteo lluvia": tablas["lluvia_conteo"],
        "Desempeño lluvia": tablas["lluvia_desempeno"],
        "Spearman": tablas["spearman"],
        "Estabilidad": tablas["estabilidad"],
    }

    with pd.ExcelWriter(XLSX_SALIDA, engine="xlsxwriter") as writer:
        workbook = writer.book
        formato_titulo = workbook.add_format(
            {"bold": True, "font_size": 14}
        )
        formato_cabecera = workbook.add_format(
            {
                "bold": True,
                "text_wrap": True,
                "align": "center",
                "valign": "vcenter",
                "border": 1,
            }
        )

        indice = workbook.add_worksheet("Índice")
        indice.write(0, 0, "Resultados de la tesis", formato_titulo)
        indice.write_row(2, 0, ["N.º", "Hoja"], formato_cabecera)

        for numero, (nombre, df) in enumerate(hojas.items(), start=1):
            df.to_excel(writer, sheet_name=nombre[:31], index=False)
            ws = writer.sheets[nombre[:31]]
            ws.freeze_panes(1, 0)

            for j, columna in enumerate(df.columns):
                valores = df[columna].map(
                    lambda x: "" if pd.isna(x) else str(x)
                )
                ancho = max(
                    len(str(columna)),
                    valores.map(len).max() if not valores.empty else 0,
                ) + 2
                ws.set_column(j, j, min(ancho, 35))

            indice.write(numero + 2, 0, numero)
            indice.write_url(
                numero + 2,
                1,
                f"internal:'{nombre[:31]}'!A1",
                string=nombre,
            )

        indice.set_column("A:A", 7)
        indice.set_column("B:B", 28)


def main() -> None:
    REPORTES_DIR.mkdir(parents=True, exist_ok=True)
    tablas = cargar_tablas()

    print("Generando Excel consolidado...")
    crear_excel(tablas)

    print("Generando Capítulo IV en Word...")
    crear_word(tablas)

    print("\nProductos generados:")
    print(f" - {XLSX_SALIDA}")
    print(f" - {DOCX_SALIDA}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"\nERROR: {error}", file=sys.stderr)
        raise
