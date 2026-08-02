"""
13_generar_capitulo_iv_final.py
==================================

Versión final del generador del Capítulo IV.

Correcciones incorporadas:
1. Calcula e incorpora Shapiro-Wilk antes de Kruskal-Wallis.
2. Recalcula la condición de lluvia con un umbral operativo configurable.
3. Genera una tabla de sensibilidad para varios umbrales de precipitación.
4. Incluye throughput en el análisis descriptivo.
5. Ordena las figuras de forma explícita.
6. Divide las tablas anchas para mejorar su legibilidad.
7. Distingue estabilidad por métrica e índice global de variabilidad.
8. Mantiene la tabla completa de Dunn en anexo.

Ejecución:
    python scripts/13_generar_capitulo_iv_final.py

Requisitos:
    python -m pip install pandas numpy scipy xlsxwriter python-docx pillow
"""

from __future__ import annotations

import hashlib
import math
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import shapiro, spearmanr
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data" / "processed"
TABLAS_DIR = BASE_DIR / "outputs" / "tables"
FIGURAS_DIR = BASE_DIR / "outputs" / "figures"
REPORTES_DIR = BASE_DIR / "outputs" / "reportes"

DATASET_ESCENARIOS = DATA_DIR / "dataset_escenarios_full_day.csv"
DATASET_CLIMA = DATA_DIR / "dataset_radio_clima.csv"

DOCX_SALIDA = REPORTES_DIR / "Capitulo_IV_Resultados_final.docx"
XLSX_SALIDA = REPORTES_DIR / "resultados_tesis_final.xlsx"

ALPHA = 0.05

# Umbral operativo para distinguir trazas de precipitación.
# Puede modificarse si la tesis adopta otro criterio.
UMBRAL_LLUVIA_MM_H = 0.10
UMBRALES_SENSIBILIDAD = [0.00, 0.01, 0.05, 0.10, 0.20, 0.50, 1.00]


DESCRIPCIONES_TABLAS = {
    "Tabla II. Resultados de Shapiro-Wilk por métrica y escenario": (
        "Con el propósito de verificar el supuesto de normalidad antes de seleccionar "
        "las pruebas inferenciales, se aplicó la prueba de Shapiro-Wilk a cada métrica "
        "y escenario experimental. La Tabla II presenta el estadístico W, el valor p "
        "y la decisión correspondiente."
    ),
    "Tabla III. Resultados de Kruskal-Wallis": (
        "Una vez confirmado el incumplimiento del supuesto de normalidad, se evaluó "
        "si las distribuciones de RSSI, SNR, MCS y throughput diferían entre los "
        "escenarios experimentales. La Tabla III resume los resultados de la prueba "
        "de Kruskal-Wallis y el tamaño del efecto."
    ),
    "Tabla IV. Síntesis de comparaciones post hoc de Dunn": (
        "Debido a que la prueba global de Kruskal-Wallis identificó diferencias "
        "significativas, se aplicaron comparaciones post hoc de Dunn con ajuste de "
        "Holm. La Tabla IV sintetiza la cantidad de comparaciones significativas y "
        "los pares que no presentaron diferencias estadísticas."
    ),
    "Tabla V. RSSI, SNR y MCS por escenario experimental": (
        "Para interpretar técnicamente las diferencias inferenciales, se examinaron "
        "los valores descriptivos de calidad de señal, calidad de canal y modulación "
        "adaptativa. La Tabla V presenta el comportamiento de RSSI, SNR y MCS en cada "
        "escenario experimental."
    ),
    "Tabla VI. Throughput observado por escenario experimental": (
        "La capacidad efectiva observada durante el monitoreo fue analizada mediante "
        "el throughput descendente. La Tabla VI presenta su media, desviación estándar "
        "y mediana por escenario, considerando que esta métrica representa tráfico "
        "cursado y no la capacidad máxima teórica del enlace."
    ),
    "Tabla VII. Sensibilidad de la clasificación de lluvia al umbral adoptado": (
        "Antes de clasificar los registros en condiciones con lluvia y sin lluvia, "
        "se examinó la sensibilidad del resultado frente a distintos umbrales de "
        "precipitación. La Tabla VII permite verificar cómo cambia la distribución "
        "de observaciones según el umbral adoptado."
    ),
    "Tabla VIII. Registros según condición de precipitación validada": (
        "A partir del umbral operativo seleccionado, los registros fueron clasificados "
        "según la presencia o ausencia de precipitación. La Tabla VIII muestra la "
        "cantidad y el porcentaje de observaciones en cada condición."
    ),
    "Tabla IX-A. RSSI y SNR según condición de precipitación": (
        "Con la finalidad de valorar posibles cambios en la calidad de señal y del "
        "canal, se compararon los estadísticos descriptivos de RSSI y SNR entre las "
        "condiciones con lluvia y sin lluvia. Los resultados se presentan en la "
        "Tabla IX-A."
    ),
    "Tabla IX-B. MCS y throughput según condición de precipitación": (
        "De manera complementaria, se examinó el comportamiento de la modulación "
        "adaptativa y del tráfico observado bajo ambas condiciones atmosféricas. "
        "La Tabla IX-B presenta los estadísticos de MCS y throughput."
    ),
    "Tabla X. Correlación horaria entre precipitación y métricas técnicas": (
        "Para analizar la asociación monotónica entre precipitación y desempeño, la "
        "telemetría fue agregada a escala horaria y correlacionada con ERA5-Land. "
        "La Tabla X presenta los coeficientes de Spearman, su significancia y el "
        "número de horas válidas."
    ),
    "Tabla XI-A. Estabilidad de RSSI y SNR por escenario": (
        "La estabilidad de la calidad de señal y del canal fue evaluada mediante la "
        "media, la desviación estándar y el coeficiente de variación. La Tabla XI-A "
        "resume estos indicadores para RSSI y SNR."
    ),
    "Tabla XI-B. Estabilidad de MCS y throughput por escenario": (
        "La consistencia de la modulación y del tráfico observado fue evaluada con "
        "indicadores equivalentes de dispersión relativa. La Tabla XI-B presenta los "
        "resultados de estabilidad correspondientes a MCS y throughput."
    ),
    "Tabla A. Comparaciones post hoc de Dunn completas": (
        "Con fines de trazabilidad y verificación, el anexo presenta el detalle de "
        "todas las comparaciones pareadas realizadas mediante la prueba de Dunn. "
        "La Tabla A incluye las medianas, el estadístico z, el valor p ajustado y la "
        "decisión para cada par de escenarios."
    ),
}

DESCRIPCIONES_FIGURAS = {
    "Fig. 7. Coeficiente de variación del RSSI descendente por escenario experimental.": (
        "La siguiente figura permite comparar la variabilidad relativa del RSSI "
        "entre los escenarios. Un coeficiente de variación menor representa una "
        "señal recibida más uniforme durante el periodo de evaluación."
    ),
    "Fig. 8. Distribución del RSSI descendente por escenario experimental.": (
        "Para complementar los valores promedio, la siguiente figura muestra la "
        "distribución del RSSI mediante diagramas de caja, permitiendo identificar "
        "medianas, dispersión y valores atípicos en cada escenario."
    ),
    "Fig. 9. Coeficiente de variación del MCS descendente por escenario experimental.": (
        "La siguiente figura compara la variabilidad relativa del MCS. Los valores "
        "más bajos indican una mayor permanencia del enlace en niveles de modulación "
        "consistentes."
    ),
    "Fig. 10. Coeficiente de variación del throughput descendente observado.": (
        "La siguiente figura presenta la variabilidad relativa del throughput "
        "observado. Su interpretación debe considerar que el tráfico cursado puede "
        "variar por la demanda y no únicamente por la capacidad técnica del enlace."
    ),
    "Fig. 11. Índice relativo de variabilidad operativa por escenario experimental.": (
        "Con el propósito de obtener una síntesis comparativa de la estabilidad, la "
        "siguiente figura muestra el índice relativo de variabilidad operativa. "
        "Los valores menores representan un comportamiento global más estable."
    ),
}


def exigir_archivo(ruta: Path) -> None:
    if not ruta.exists():
        raise FileNotFoundError(f"No se encontró el archivo requerido: {ruta}")


def leer_csv(ruta: Path) -> pd.DataFrame:
    exigir_archivo(ruta)
    return pd.read_csv(ruta, encoding="utf-8-sig")


def guardar_csv(df: pd.DataFrame, nombre: str) -> Path:
    TABLAS_DIR.mkdir(parents=True, exist_ok=True)
    ruta = TABLAS_DIR / nombre
    df.to_csv(ruta, index=False, encoding="utf-8-sig")
    return ruta


def formato_p(p: float) -> str:
    if pd.isna(p):
        return ""
    return "< 0.001" if p < 0.001 else f"{p:.4f}"


def calcular_shapiro(df: pd.DataFrame) -> pd.DataFrame:
    configuracion = {
        "rssi_dl": ("RSSI DL", ["E0", "E1", "E2", "E3", "E4", "E5"]),
        "snr_dl": ("SNR DL", ["E0", "E1", "E2", "E3", "E4", "E5"]),
        "mcs_dl": ("MCS DL", ["E0", "E1", "E2", "E3", "E4", "E5"]),
        "throughput_dl": ("Throughput DL", ["E1", "E2", "E3", "E4", "E5"]),
    }

    filas = []
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        for columna, (nombre, escenarios) in configuracion.items():
            for escenario in escenarios:
                valores = pd.to_numeric(
                    df.loc[df["escenario"] == escenario, columna],
                    errors="coerce",
                ).dropna()

                if len(valores) < 3:
                    continue

                w, p = shapiro(valores.to_numpy())
                filas.append(
                    {
                        "Métrica": nombre,
                        "Escenario": escenario,
                        "N": len(valores),
                        "W": round(float(w), 3),
                        "Valor p": formato_p(float(p)),
                        "Decisión": "No normal" if p < ALPHA else "Normal",
                    }
                )

    return pd.DataFrame(filas)


def generar_sensibilidad_lluvia(df: pd.DataFrame) -> pd.DataFrame:
    precip = pd.to_numeric(df["precip_mm"], errors="coerce")
    validos = precip.notna().sum()

    filas = []
    for umbral in UMBRALES_SENSIBILIDAD:
        con_lluvia = int((precip >= umbral).sum()) if umbral > 0 else int((precip > 0).sum())
        filas.append(
            {
                "Umbral (mm/h)": umbral,
                "Con lluvia (N)": con_lluvia,
                "Con lluvia (%)": round(100 * con_lluvia / validos, 2) if validos else np.nan,
                "Sin lluvia (N)": int(validos - con_lluvia),
                "Sin lluvia (%)": round(100 * (validos - con_lluvia) / validos, 2)
                if validos
                else np.nan,
            }
        )
    return pd.DataFrame(filas)


def reclasificar_lluvia(df: pd.DataFrame) -> pd.DataFrame:
    salida = df.copy()
    salida["precip_mm"] = pd.to_numeric(salida["precip_mm"], errors="coerce")
    salida["condicion_lluvia_validada"] = np.where(
        salida["precip_mm"].isna(),
        "Sin dato",
        np.where(
            salida["precip_mm"] >= UMBRAL_LLUVIA_MM_H,
            "Con lluvia",
            "Sin lluvia",
        ),
    )
    return salida


def tabla_conteo_lluvia(df: pd.DataFrame) -> pd.DataFrame:
    validos = df[df["condicion_lluvia_validada"] != "Sin dato"]
    conteo = (
        validos["condicion_lluvia_validada"]
        .value_counts()
        .rename_axis("Condición")
        .reset_index(name="N")
    )
    conteo["Porcentaje (%)"] = (100 * conteo["N"] / len(validos)).round(2)
    orden = pd.Categorical(
        conteo["Condición"],
        categories=["Con lluvia", "Sin lluvia"],
        ordered=True,
    )
    return conteo.assign(_orden=orden).sort_values("_orden").drop(columns="_orden")


def tabla_desempeno_lluvia(df: pd.DataFrame) -> pd.DataFrame:
    metricas = ["rssi_dl", "snr_dl", "mcs_dl", "throughput_dl"]
    etiquetas = {
        "rssi_dl": "RSSI",
        "snr_dl": "SNR",
        "mcs_dl": "MCS",
        "throughput_dl": "Throughput",
    }

    filas = []
    for condicion in ["Con lluvia", "Sin lluvia"]:
        grupo = df[df["condicion_lluvia_validada"] == condicion]
        fila = {"Condición": condicion}
        for metrica in metricas:
            valores = pd.to_numeric(grupo[metrica], errors="coerce").dropna()
            fila[f"{etiquetas[metrica]} media"] = valores.mean()
            fila[f"{etiquetas[metrica]} DE"] = valores.std(ddof=1)
            fila[f"{etiquetas[metrica]} mediana"] = valores.median()
        filas.append(fila)

    return pd.DataFrame(filas).round(3)


def tabla_spearman_horaria(df: pd.DataFrame) -> pd.DataFrame:
    """
    Agrega la telemetría por hora antes de correlacionarla con precipitación,
    evitando repetir el mismo valor horario de ERA5-Land en cada registro minuto.
    """
    if "hora_merge" not in df.columns:
        raise ValueError("dataset_radio_clima.csv no contiene 'hora_merge'.")

    columnas = ["precip_mm", "rssi_dl", "snr_dl", "mcs_dl", "throughput_dl"]
    horario = (
        df.groupby("hora_merge", as_index=False)[columnas]
        .mean(numeric_only=True)
    )

    etiquetas = {
        "rssi_dl": "RSSI DL",
        "snr_dl": "SNR DL",
        "mcs_dl": "MCS DL",
        "throughput_dl": "Throughput DL",
    }

    filas = []
    for metrica, nombre in etiquetas.items():
        pares = horario[["precip_mm", metrica]].dropna()
        if len(pares) < 3:
            rho, p = np.nan, np.nan
        else:
            rho, p = spearmanr(pares["precip_mm"], pares[metrica])

        filas.append(
            {
                "Métrica": nombre,
                "ρ de Spearman": round(float(rho), 3) if pd.notna(rho) else np.nan,
                "Valor p": formato_p(float(p)) if pd.notna(p) else "",
                "N horas": len(pares),
            }
        )

    return pd.DataFrame(filas)


def tabla_descriptivos_con_throughput(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    base = (
        df.groupby(["escenario", "frecuencia_mhz", "ancho_canal_mhz"], as_index=False)
        .agg(
            N=("escenario", "size"),
            Dias=("fecha", "nunique"),
            RSSI_media=("rssi_dl", "mean"),
            RSSI_DE=("rssi_dl", "std"),
            SNR_media=("snr_dl", "mean"),
            SNR_DE=("snr_dl", "std"),
            MCS_media=("mcs_dl", "mean"),
            MCS_DE=("mcs_dl", "std"),
            Throughput_media=("throughput_dl", "mean"),
            Throughput_DE=("throughput_dl", "std"),
            Throughput_mediana=("throughput_dl", "median"),
        )
    )

    senal = base[
        [
            "escenario",
            "frecuencia_mhz",
            "ancho_canal_mhz",
            "N",
            "Dias",
            "RSSI_media",
            "RSSI_DE",
            "SNR_media",
            "SNR_DE",
            "MCS_media",
            "MCS_DE",
        ]
    ].copy()

    capacidad = base[
        [
            "escenario",
            "frecuencia_mhz",
            "ancho_canal_mhz",
            "Throughput_media",
            "Throughput_DE",
            "Throughput_mediana",
        ]
    ].copy()

    renombres = {
        "escenario": "Escenario",
        "frecuencia_mhz": "Frecuencia (MHz)",
        "ancho_canal_mhz": "Canal (MHz)",
        "Dias": "Días",
        "RSSI_media": "RSSI media (dBm)",
        "RSSI_DE": "RSSI DE",
        "SNR_media": "SNR media (dB)",
        "SNR_DE": "SNR DE",
        "MCS_media": "MCS media",
        "MCS_DE": "MCS DE",
        "Throughput_media": "Throughput media (Mbps)",
        "Throughput_DE": "Throughput DE",
        "Throughput_mediana": "Throughput mediana (Mbps)",
    }

    return senal.rename(columns=renombres).round(3), capacidad.rename(columns=renombres).round(3)


def cargar_csv_existente(nombre: str) -> pd.DataFrame:
    return leer_csv(TABLAS_DIR / nombre)


def resumen_dunn(completa: pd.DataFrame) -> pd.DataFrame:
    filas = []
    for metrica, grupo in completa.groupby("Métrica", sort=False):
        sig = grupo["Diferencia significativa"].astype(str).str.strip().eq("Sí")
        no_sig = grupo.loc[~sig]
        pares = (
            "Ninguna"
            if no_sig.empty
            else "; ".join(
                f"{r['Comparación']} (p={r['Valor p ajustado']})"
                for _, r in no_sig.iterrows()
            )
        )
        filas.append(
            {
                "Métrica": metrica,
                "Significativas": f"{int(sig.sum())} de {len(grupo)}",
                "No significativas": pares,
            }
        )
    return pd.DataFrame(filas)


def natural_key(texto: str) -> list[object]:
    return [
        int(fragmento) if fragmento.isdigit() else fragmento.lower()
        for fragmento in re.split(r"(\d+)", texto)
    ]


FIGURAS_ORDEN = [
    ("fig_4_9_cv_rssi", "Fig. 7. Coeficiente de variación del RSSI descendente por escenario experimental."),
    ("fig_4_10_boxplot_rssi", "Fig. 8. Distribución del RSSI descendente por escenario experimental."),
    ("fig_4_11_cv_mcs", "Fig. 9. Coeficiente de variación del MCS descendente por escenario experimental."),
    ("fig_4_11_cv_throughput", "Fig. 10. Coeficiente de variación del throughput descendente observado."),
    ("fig_4_12_indice_variabilidad", "Fig. 11. Índice relativo de variabilidad operativa por escenario experimental."),
]


def figuras_ordenadas() -> list[tuple[Path, str]]:
    resultado = []
    hashes = set()

    for stem, titulo in FIGURAS_ORDEN:
        candidatas = sorted(FIGURAS_DIR.glob(f"{stem}.*"), key=lambda p: natural_key(p.name))
        for ruta in candidatas:
            if ruta.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                continue
            digest = hashlib.sha256(ruta.read_bytes()).hexdigest()
            if digest in hashes:
                continue
            hashes.add(digest)
            resultado.append((ruta, titulo))
            break

    return resultado


def configurar_documento() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    return doc


def bordes(celda) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "B7B7B7")
        tc_borders.append(e)
    tc_pr.append(tc_borders)


def sombrear(celda, color: str = "D9EAF7") -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(shd)


def parrafo(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(texto)
    r.font.name = "Arial"
    r.font.size = Pt(10)


def tabla_word(
    doc: Document,
    titulo: str,
    df: pd.DataFrame,
    nota: str | None = None,
) -> None:
    descripcion = DESCRIPCIONES_TABLAS.get(titulo)
    if descripcion:
        parrafo(doc, descripcion)

    p = doc.add_paragraph()
    r = p.add_run(titulo)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)

    tabla = doc.add_table(rows=1, cols=len(df.columns))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = True

    for i, col in enumerate(df.columns):
        celda = tabla.rows[0].cells[i]
        celda.text = str(col)
        sombrear(celda)
        bordes(celda)
        for pp in celda.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for rr in pp.runs:
                rr.bold = True
                rr.font.name = "Arial"
                rr.font.size = Pt(8)

    for _, fila in df.iterrows():
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            if pd.isna(valor):
                texto = ""
            elif isinstance(valor, (float, np.floating)):
                texto = f"{valor:.3f}".rstrip("0").rstrip(".")
            else:
                texto = str(valor)
            celdas[i].text = texto
            bordes(celdas[i])
            celdas[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for pp in celdas[i].paragraphs:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs:
                    rr.font.name = "Arial"
                    rr.font.size = Pt(8)

    if nota:
        p = doc.add_paragraph()
        r = p.add_run(nota)
        r.italic = True
        r.font.name = "Arial"
        r.font.size = Pt(8)

    doc.add_paragraph()


def agregar_figura(doc: Document, ruta: Path, titulo: str) -> None:
    descripcion = DESCRIPCIONES_FIGURAS.get(titulo)
    if descripcion:
        parrafo(doc, descripcion)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=Cm(15.5))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fuente: Elaboración propia.")
    r.italic = True
    r.font.size = Pt(8)


def obtener_indice_global() -> tuple[str | None, float | None]:
    ruta = FIGURAS_DIR / "tabla_estabilidad_calculada.csv"
    if not ruta.exists():
        return None, None

    df = pd.read_csv(ruta)
    columnas_lower = {c.lower(): c for c in df.columns}
    col_esc = next((c for k, c in columnas_lower.items() if "escenario" in k), None)
    col_ind = next(
        (
            c
            for k, c in columnas_lower.items()
            if "indice" in k or "índice" in k or "variabilidad" in k
        ),
        None,
    )
    if not col_esc or not col_ind:
        return None, None

    valores = pd.to_numeric(df[col_ind], errors="coerce")
    if valores.notna().sum() == 0:
        return None, None

    idx = valores.idxmin()
    return str(df.loc[idx, col_esc]), float(valores.loc[idx])


def crear_word(tablas: dict[str, pd.DataFrame]) -> None:
    doc = configurar_documento()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CAPÍTULO IV\nRESULTADOS")
    r.bold = True
    r.font.size = Pt(14)

    doc.add_heading("4.1. Verificación de normalidad y selección de pruebas", level=1)
    parrafo(
        doc,
        "Antes de seleccionar las pruebas inferenciales, se evaluó la normalidad "
        "de RSSI, SNR, MCS y throughput mediante Shapiro-Wilk, con α=0.05. "
        "En todas las combinaciones de métrica y escenario se obtuvo p<0.001; "
        "por tanto, se rechazó la hipótesis de normalidad y se seleccionaron "
        "procedimientos no paramétricos."
    )
    tabla_word(
        doc,
        "Tabla II. Resultados de Shapiro-Wilk por métrica y escenario",
        tablas["shapiro"],
        nota=(
            "Nota. Para E0, N superó 5000 observaciones; el resultado se interpretó "
            "junto con la naturaleza discreta o asimétrica de las distribuciones."
        ),
    )

    doc.add_heading("4.2. Contraste inferencial entre escenarios", level=1)
    tabla_word(
        doc,
        "Tabla III. Resultados de Kruskal-Wallis",
        tablas["kruskal"],
        nota=(
            "Nota. E0 se excluyó del contraste de throughput por no disponer de "
            "registros equivalentes a cnMaestro."
        ),
    )
    parrafo(
        doc,
        "Kruskal-Wallis evidenció diferencias significativas para las cuatro "
        "métricas (p<0.001), con tamaños del efecto grandes. En consecuencia, "
        "la configuración técnica influyó sobre el desempeño del radioenlace."
    )
    tabla_word(
        doc,
        "Tabla IV. Síntesis de comparaciones post hoc de Dunn",
        tablas["dunn_resumen"],
        nota="Nota. Los valores p fueron ajustados mediante Holm.",
    )

    doc.add_heading("4.3. Comparación descriptiva entre escenarios", level=1)
    tabla_word(
        doc,
        "Tabla V. RSSI, SNR y MCS por escenario experimental",
        tablas["descriptivos_senal"],
    )
    tabla_word(
        doc,
        "Tabla VI. Throughput observado por escenario experimental",
        tablas["descriptivos_throughput"],
        nota=(
            "Nota. El throughput representa tráfico observado durante el monitoreo "
            "y no la capacidad máxima teórica del enlace."
        ),
    )

    th = tablas["descriptivos_throughput"].dropna(
        subset=["Throughput mediana (Mbps)"]
    )
    if not th.empty:
        mejor = th.loc[th["Throughput mediana (Mbps)"].idxmax()]
        parrafo(
            doc,
            f"El escenario {mejor['Escenario']} presentó la mayor mediana de "
            f"throughput observado ({mejor['Throughput mediana (Mbps)']:.2f} Mbps)."
        )

    doc.add_heading("4.4. Precipitación como covariable", level=1)
    tabla_word(
        doc,
        "Tabla VII. Sensibilidad de la clasificación de lluvia al umbral adoptado",
        tablas["sensibilidad_lluvia"],
        nota=(
            f"Nota. Para el análisis principal se utilizó un umbral operativo de "
            f"{UMBRAL_LLUVIA_MM_H:.2f} mm/h."
        ),
    )
    tabla_word(
        doc,
        "Tabla VIII. Registros según condición de precipitación validada",
        tablas["conteo_lluvia"],
    )
    tabla_word(
        doc,
        "Tabla IX-A. RSSI y SNR según condición de precipitación",
        tablas["lluvia_senal"],
    )
    tabla_word(
        doc,
        "Tabla IX-B. MCS y throughput según condición de precipitación",
        tablas["lluvia_capacidad"],
    )
    tabla_word(
        doc,
        "Tabla X. Correlación horaria entre precipitación y métricas técnicas",
        tablas["spearman"],
        nota=(
            "Nota. La telemetría se agregó por hora antes de correlacionarla con "
            "ERA5-Land, evitando repetir el mismo valor de precipitación."
        ),
    )

    doc.add_heading("4.5. Estabilidad operativa", level=1)
    tabla_word(
        doc,
        "Tabla XI-A. Estabilidad de RSSI y SNR por escenario",
        tablas["estabilidad_senal"],
    )
    tabla_word(
        doc,
        "Tabla XI-B. Estabilidad de MCS y throughput por escenario",
        tablas["estabilidad_capacidad"],
    )

    escenario, indice = obtener_indice_global()
    if escenario is not None:
        parrafo(
            doc,
            f"Aunque la estabilidad relativa varió según la métrica, el índice "
            f"compuesto identificó a {escenario} como el escenario globalmente "
            f"más estable, con un valor normalizado de {indice:.2f}."
        )
    else:
        parrafo(
            doc,
            "La estabilidad relativa varió según la métrica evaluada. El índice "
            "compuesto debe interpretarse como una síntesis complementaria y no "
            "como sustituto de los indicadores individuales."
        )

    parrafo(
        doc,
        "Las figuras siguientes complementan las tablas de estabilidad y permiten "
        "visualizar la dispersión, la presencia de valores atípicos y la posición "
        "relativa de cada escenario respecto de la variabilidad operativa."
    )

    for ruta, titulo in figuras_ordenadas():
        agregar_figura(doc, ruta, titulo)
        if "Índice relativo" in titulo:
            parrafo(
                doc,
                "Valores menores del índice representan menor variabilidad "
                "operativa y, por tanto, mayor estabilidad relativa."
            )

    doc.add_page_break()
    doc.add_heading("ANEXO. Comparaciones post hoc de Dunn completas", level=1)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.3)
    sec.right_margin = Cm(1.3)

    tabla_word(
        doc,
        "Tabla A. Comparaciones post hoc de Dunn completas",
        tablas["dunn_completa"],
        nota="Nota. Los valores p fueron ajustados mediante Holm.",
    )

    doc.save(DOCX_SALIDA)


def crear_excel(tablas: dict[str, pd.DataFrame]) -> None:
    hojas = {
        "Shapiro": tablas["shapiro"],
        "Kruskal-Wallis": tablas["kruskal"],
        "Dunn resumen": tablas["dunn_resumen"],
        "Dunn completa": tablas["dunn_completa"],
        "Descriptivos señal": tablas["descriptivos_senal"],
        "Descriptivos throughput": tablas["descriptivos_throughput"],
        "Sensibilidad lluvia": tablas["sensibilidad_lluvia"],
        "Conteo lluvia": tablas["conteo_lluvia"],
        "Lluvia señal": tablas["lluvia_senal"],
        "Lluvia capacidad": tablas["lluvia_capacidad"],
        "Spearman horario": tablas["spearman"],
        "Estabilidad señal": tablas["estabilidad_senal"],
        "Estabilidad capacidad": tablas["estabilidad_capacidad"],
    }

    with pd.ExcelWriter(XLSX_SALIDA, engine="xlsxwriter") as writer:
        indice = writer.book.add_worksheet("Índice")
        indice.write_row(0, 0, ["N.º", "Hoja"])

        for numero, (nombre, df) in enumerate(hojas.items(), start=1):
            hoja = nombre[:31]
            df.to_excel(writer, sheet_name=hoja, index=False)
            ws = writer.sheets[hoja]
            ws.freeze_panes(1, 0)

            for j, col in enumerate(df.columns):
                serie = df[col].map(lambda x: "" if pd.isna(x) else str(x))
                max_len = int(serie.map(len).max()) if not serie.empty else 0
                ws.set_column(j, j, min(max(len(str(col)), max_len) + 2, 35))

            indice.write(numero, 0, numero)
            indice.write_url(numero, 1, f"internal:'{hoja}'!A1", string=hoja)

        indice.set_column("A:A", 7)
        indice.set_column("B:B", 28)


def main() -> None:
    REPORTES_DIR.mkdir(parents=True, exist_ok=True)
    TABLAS_DIR.mkdir(parents=True, exist_ok=True)

    escenarios = leer_csv(DATASET_ESCENARIOS)
    clima = leer_csv(DATASET_CLIMA)

    shapiro_df = calcular_shapiro(escenarios)
    guardar_csv(shapiro_df, "tabla_4_1_normalidad_shapiro.csv")

    sensibilidad = generar_sensibilidad_lluvia(clima)
    clima_validado = reclasificar_lluvia(clima)
    conteo = tabla_conteo_lluvia(clima_validado)
    desempeno_lluvia = tabla_desempeno_lluvia(clima_validado)
    spearman_h = tabla_spearman_horaria(clima_validado)

    guardar_csv(sensibilidad, "tabla_4_7_sensibilidad_umbral_lluvia.csv")
    guardar_csv(conteo, "tabla_4_8_conteo_lluvia_validado.csv")
    guardar_csv(desempeno_lluvia, "tabla_4_9_desempeno_lluvia_validado.csv")
    guardar_csv(spearman_h, "tabla_4_10_spearman_horario.csv")

    desc_senal, desc_th = tabla_descriptivos_con_throughput(escenarios)

    kruskal = cargar_csv_existente("tabla_4_2_kruskal_wallis.csv")
    dunn_completa = cargar_csv_existente("tabla_4_3_dunn_completa.csv")
    estabilidad = cargar_csv_existente("tabla_4_6_estabilidad_por_escenario.csv")

    estabilidad = estabilidad.rename(
        columns={
            "escenario": "Escenario",
            "rssi_dl_mean": "RSSI media (dBm)",
            "rssi_dl_std": "RSSI DE",
            "rssi_dl_cv": "RSSI CV (%)",
            "snr_dl_mean": "SNR media (dB)",
            "snr_dl_std": "SNR DE",
            "snr_dl_cv": "SNR CV (%)",
            "mcs_dl_mean": "MCS media",
            "mcs_dl_std": "MCS DE",
            "mcs_dl_cv": "MCS CV (%)",
            "throughput_dl_mean": "Throughput media (Mbps)",
            "throughput_dl_std": "Throughput DE",
            "throughput_dl_cv": "Throughput CV (%)",
        }
    ).round(3)

    estabilidad_senal = estabilidad[
        [
            "Escenario",
            "RSSI media (dBm)",
            "RSSI DE",
            "RSSI CV (%)",
            "SNR media (dB)",
            "SNR DE",
            "SNR CV (%)",
        ]
    ]
    estabilidad_capacidad = estabilidad[
        [
            "Escenario",
            "MCS media",
            "MCS DE",
            "MCS CV (%)",
            "Throughput media (Mbps)",
            "Throughput DE",
            "Throughput CV (%)",
        ]
    ]

    lluvia_senal = desempeno_lluvia[
        [
            "Condición",
            "RSSI media",
            "RSSI DE",
            "RSSI mediana",
            "SNR media",
            "SNR DE",
            "SNR mediana",
        ]
    ]
    lluvia_capacidad = desempeno_lluvia[
        [
            "Condición",
            "MCS media",
            "MCS DE",
            "MCS mediana",
            "Throughput media",
            "Throughput DE",
            "Throughput mediana",
        ]
    ]

    tablas = {
        "shapiro": shapiro_df,
        "kruskal": kruskal,
        "dunn_resumen": resumen_dunn(dunn_completa),
        "dunn_completa": dunn_completa,
        "descriptivos_senal": desc_senal,
        "descriptivos_throughput": desc_th,
        "sensibilidad_lluvia": sensibilidad,
        "conteo_lluvia": conteo,
        "lluvia_senal": lluvia_senal,
        "lluvia_capacidad": lluvia_capacidad,
        "spearman": spearman_h,
        "estabilidad_senal": estabilidad_senal,
        "estabilidad_capacidad": estabilidad_capacidad,
    }

    print("Generando Excel validado...")
    crear_excel(tablas)

    print("Generando Capítulo IV validado...")
    crear_word(tablas)

    print("\nProductos generados:")
    print(f" - {XLSX_SALIDA}")
    print(f" - {DOCX_SALIDA}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"\nERROR: {error}", file=sys.stderr)
        raise
