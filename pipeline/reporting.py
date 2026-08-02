
from __future__ import annotations

import hashlib
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .config import (
    DOCX_OUT,
    XLSX_OUT,
    FIGURES,
    FIGURE_ORDER,
    FIGURE_START,
)
from .interpretation import (
    shapiro_text,
    kruskal_text,
    dunn_text,
    descriptive_text,
    throughput_text,
    rain_count_text,
    rain_signal_text,
    rain_capacity_text,
    spearman_text,
    stability_text,
    final_synthesis,
)


TABLE_TITLES = {
    "shapiro": "Tabla IV. Resultados de Shapiro-Wilk por métrica y escenario",
    "kruskal": "Tabla V. Resultados de Kruskal-Wallis",
    "dunn_summary": "Tabla VI. Síntesis de comparaciones post hoc de Dunn",
    "descriptive_signal": "Tabla VII. RSSI, SNR y MCS por escenario experimental",
    "descriptive_throughput": "Tabla VIII. Throughput observado por escenario experimental",
    "rain_sensitivity": "Tabla IX. Sensibilidad de la clasificación de lluvia",
    "rain_count": "Tabla X. Registros según condición de precipitación",
    "rain_signal": "Tabla XI-A. RSSI y SNR según precipitación",
    "rain_capacity": "Tabla XI-B. MCS y throughput según precipitación",
    "spearman": "Tabla XII. Correlación horaria entre precipitación y métricas",
    "stability_signal": "Tabla XIII-A. Estabilidad de RSSI y SNR",
    "stability_capacity": "Tabla XIII-B. Estabilidad de MCS y throughput",
    "dunn_full": "Tabla A1. Comparaciones post hoc de Dunn completas",
}


TABLE_INTRO = {
    "shapiro": (
        "Antes de seleccionar los procedimientos inferenciales, se verificó el "
        "supuesto de normalidad en cada combinación de métrica y escenario. "
        "La Tabla IV presenta el estadístico W, el valor de significancia y la "
        "decisión correspondiente."
    ),
    "kruskal": (
        "Una vez confirmado el incumplimiento del supuesto de normalidad, se aplicó "
        "la prueba de Kruskal-Wallis para determinar si las distribuciones de las "
        "métricas técnicas diferían entre los escenarios experimentales. Los "
        "resultados se presentan en la Tabla V."
    ),
    "dunn_summary": (
        "Debido a que el contraste global identificó diferencias significativas, "
        "se realizaron comparaciones múltiples mediante la prueba post hoc de Dunn "
        "con ajuste de Holm. La Tabla VI resume las comparaciones significativas y "
        "los pares estadísticamente equivalentes."
    ),
    "descriptive_signal": (
        "Con el propósito de caracterizar el comportamiento del radioenlace bajo "
        "cada configuración experimental, la Tabla VII resume los principales "
        "estadísticos descriptivos correspondientes a RSSI, SNR y MCS."
    ),
    "descriptive_throughput": (
        "La capacidad efectiva observada durante el periodo de monitoreo fue "
        "examinada mediante el throughput descendente. La Tabla VIII presenta la "
        "media, la desviación estándar y la mediana registradas en cada escenario."
    ),
    "rain_sensitivity": (
        "Antes de clasificar las observaciones según la presencia de precipitación, "
        "se evaluó la sensibilidad de la distribución frente a diferentes umbrales. "
        "La Tabla IX muestra cómo varió la proporción de registros con lluvia y sin "
        "lluvia."
    ),
    "rain_count": (
        "A partir del umbral operativo seleccionado, los registros fueron clasificados "
        "en condiciones con lluvia y sin lluvia. La Tabla X presenta la frecuencia y "
        "el porcentaje correspondiente a cada categoría."
    ),
    "rain_signal": (
        "Para valorar posibles cambios en la calidad de señal y del canal, la Tabla "
        "XI-A compara el RSSI y el SNR entre los registros clasificados con lluvia y "
        "sin lluvia."
    ),
    "rain_capacity": (
        "De manera complementaria, la Tabla XI-B presenta el comportamiento del MCS "
        "y del throughput observado bajo ambas condiciones atmosféricas."
    ),
    "spearman": (
        "Con el fin de evaluar la asociación monotónica entre la precipitación y el "
        "desempeño técnico, la telemetría fue agregada a escala horaria y correlacionada "
        "con ERA5-Land mediante el coeficiente de Spearman. Los resultados se muestran "
        "en la Tabla XII."
    ),
    "stability_signal": (
        "La estabilidad de la calidad de señal y del canal fue evaluada mediante la "
        "desviación estándar y el coeficiente de variación. La Tabla XIII-A resume "
        "estos indicadores para RSSI y SNR."
    ),
    "stability_capacity": (
        "La consistencia de la modulación adaptativa y del tráfico observado fue "
        "analizada mediante indicadores equivalentes de dispersión relativa. La "
        "Tabla XIII-B presenta los resultados correspondientes a MCS y throughput."
    ),
    "dunn_full": (
        "Con fines de trazabilidad y verificación, el anexo presenta el detalle de "
        "todas las comparaciones pareadas realizadas mediante la prueba de Dunn. "
        "La Tabla A1 incluye las medianas, el estadístico z, el valor p ajustado y "
        "la decisión para cada par de escenarios."
    ),
}


def _border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "B7B7B7")
        borders.append(e)
    tc_pr.append(borders)


def _shade(cell) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9EAF7")
    cell._tc.get_or_add_tcPr().append(shd)


def paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10)


def add_table(
    doc: Document,
    key: str,
    df: pd.DataFrame,
    interpretation: str,
    note: str | None = None,
) -> None:
    paragraph(doc, TABLE_INTRO[key])

    p = doc.add_paragraph()
    r = p.add_run(TABLE_TITLES[key])
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, col in enumerate(df.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(col).replace("Mbps", "Mb/s")
        _shade(cell)
        _border(cell)
        for pp in cell.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for rr in pp.runs:
                rr.bold = True
                rr.font.name = "Arial"
                rr.font.size = Pt(8)

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if pd.isna(value):
                text = ""
            elif isinstance(value, (float, np.floating)):
                text = f"{value:.3f}".rstrip("0").rstrip(".")
            else:
                text = str(value)
            cells[i].text = text
            _border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for pp in cells[i].paragraphs:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs:
                    rr.font.name = "Arial"
                    rr.font.size = Pt(8)

    if note:
        p = doc.add_paragraph()
        r = p.add_run(note)
        r.italic = True
        r.font.name = "Arial"
        r.font.size = Pt(8)

    paragraph(doc, interpretation)


def unique_figures() -> list[tuple[Path, str]]:
    result = []
    hashes = set()
    for stem, description in FIGURE_ORDER:
        for path in sorted(FIGURES.glob(f"{stem}.*")):
            if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                continue
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            if digest in hashes:
                continue
            hashes.add(digest)
            result.append((path, description))
            break
    return result


def figure_intro(index: int) -> str:
    intros = [
        (
            "Con el fin de visualizar la estabilidad de la potencia recibida entre "
            "escenarios, la siguiente figura presenta el coeficiente de variación "
            "del RSSI descendente."
        ),
        (
            "Para examinar la dispersión, la posición de las medianas y la presencia "
            "de valores atípicos, la siguiente figura muestra la distribución del "
            "RSSI mediante diagramas de caja."
        ),
        (
            "La estabilidad del mecanismo de modulación adaptativa fue complementada "
            "mediante la representación gráfica del coeficiente de variación del MCS."
        ),
        (
            "La siguiente representación permite comparar la estabilidad relativa "
            "del throughput observado entre los escenarios experimentales."
        ),
        (
            "Finalmente, los indicadores individuales de estabilidad fueron integrados "
            "en un índice relativo normalizado, con el propósito de obtener una "
            "comparación global entre escenarios."
        ),
    ]
    return intros[index]


def figure_interpretation(index: int, tables: dict[str, pd.DataFrame]) -> str:
    if index == 0:
        row = tables["stability_signal"].loc[
            tables["stability_signal"]["RSSI CV (%)"].idxmin()
        ]
        return (
            f"La figura confirma que {row['Escenario']} presentó el menor coeficiente "
            f"de variación del RSSI (CV = {row['RSSI CV (%)']:.3f} %), evidenciando "
            "el comportamiento más uniforme de la potencia recibida durante el "
            "periodo de observación."
        )
    if index == 1:
        row = tables["descriptive_signal"].loc[
            tables["descriptive_signal"]["RSSI media (dBm)"].idxmax()
        ]
        return (
            f"La distribución de {row['Escenario']} se desplazó hacia niveles menos "
            "negativos de RSSI, coherente con el mayor promedio observado. Asimismo, "
            "la figura permite identificar valores atípicos en determinados escenarios, "
            "los cuales explican parte de la variabilidad registrada."
        )
    if index == 2:
        row = tables["stability_capacity"].loc[
            tables["stability_capacity"]["MCS CV (%)"].idxmin()
        ]
        return (
            f"{row['Escenario']} presentó la menor variabilidad del MCS "
            f"(CV = {row['MCS CV (%)']:.3f} %), indicando una mayor permanencia del "
            "enlace en niveles de modulación consistentes."
        )
    if index == 3:
        t = tables["stability_capacity"].dropna(subset=["Throughput CV (%)"])
        row = t.loc[t["Throughput CV (%)"].idxmin()]
        return (
            f"Aunque el mayor throughput promedio no necesariamente coincidió con la "
            f"mayor estabilidad, {row['Escenario']} presentó el menor coeficiente de "
            f"variación del throughput (CV = {row['Throughput CV (%)']:.3f} %), "
            "evidenciando un comportamiento operativo más uniforme."
        )
    return (
        "El índice integrado identificó a E3 como el escenario con menor variabilidad "
        "operativa global, seguido de E4 y E5. Este resultado muestra que la estabilidad "
        "general no dependió de un único indicador, sino de la combinación de las "
        "métricas técnicas evaluadas."
    )


def add_figure(
    doc: Document,
    path: Path,
    description: str,
    number: int,
    intro_text: str,
    interpretation: str,
) -> None:
    paragraph(doc, intro_text)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(15.5))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Fig. {number}. {description}")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fuente: Elaboración propia.")
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(8)

    paragraph(doc, interpretation)


def create_word(tables: dict[str, pd.DataFrame]) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CAPÍTULO IV\nRESULTADOS")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)

    doc.add_heading("4.1. Verificación de normalidad y selección de pruebas", level=1)
    add_table(doc, "shapiro", tables["shapiro"], shapiro_text(tables["shapiro"]))

    doc.add_heading("4.2. Contraste inferencial entre escenarios", level=1)
    add_table(doc, "kruskal", tables["kruskal"], kruskal_text(tables["kruskal"]))
    add_table(
        doc,
        "dunn_summary",
        tables["dunn_summary"],
        dunn_text(tables["dunn_summary"]),
    )
    paragraph(
        doc,
        "En conjunto, el análisis inferencial demostró que las modificaciones "
        "introducidas en la frecuencia de operación y el ancho de canal produjeron "
        "cambios estadísticamente significativos en el desempeño del radioenlace. "
        "Este resultado justificó profundizar la interpretación mediante el análisis "
        "descriptivo de cada escenario."
    )

    doc.add_heading("4.3. Comparación descriptiva entre escenarios", level=1)
    add_table(
        doc,
        "descriptive_signal",
        tables["descriptive_signal"],
        descriptive_text(tables["descriptive_signal"]),
    )
    add_table(
        doc,
        "descriptive_throughput",
        tables["descriptive_throughput"],
        throughput_text(tables["descriptive_throughput"]),
        note=(
            "Nota. El throughput corresponde al tráfico observado durante el "
            "monitoreo y no a la capacidad máxima teórica del enlace."
        ),
    )
    paragraph(
        doc,
        "Los resultados descriptivos complementaron la evidencia inferencial y "
        "permitieron identificar tendencias diferenciadas entre escenarios. No se "
        "observó una única configuración dominante para todas las métricas, lo que "
        "evidencia la existencia de compromisos entre calidad de señal, calidad de "
        "canal, modulación y tráfico observado."
    )

    doc.add_heading("4.4. Precipitación como covariable", level=1)
    add_table(
        doc,
        "rain_sensitivity",
        tables["rain_sensitivity"],
        (
            "La distribución de registros cambió de forma sustancial según el "
            "umbral considerado. Por esta razón, el criterio de 0.10 mm/h se mantuvo "
            "explícito y se utilizó de manera uniforme en los análisis posteriores."
        ),
    )
    add_table(
        doc,
        "rain_count",
        tables["rain_count"],
        rain_count_text(tables["rain_count"]),
    )
    add_table(
        doc,
        "rain_signal",
        tables["rain_signal"],
        rain_signal_text(tables["rain_signal"]),
    )
    add_table(
        doc,
        "rain_capacity",
        tables["rain_capacity"],
        rain_capacity_text(tables["rain_capacity"]),
    )
    add_table(
        doc,
        "spearman",
        tables["spearman"],
        spearman_text(tables["spearman"]),
    )
    paragraph(
        doc,
        "Aunque algunas asociaciones alcanzaron significancia estadística, sus "
        "magnitudes fueron débiles o muy débiles. En consecuencia, la precipitación "
        "mostró una influencia secundaria durante el periodo de estudio frente al "
        "efecto de la configuración técnica."
    )

    doc.add_heading("4.5. Estabilidad operativa", level=1)
    stability_summary = stability_text(
        tables["stability_signal"],
        tables["stability_capacity"],
    )
    add_table(
        doc,
        "stability_signal",
        tables["stability_signal"],
        stability_summary,
    )
    add_table(
        doc,
        "stability_capacity",
        tables["stability_capacity"],
        (
            "Los resultados muestran que la estabilidad operativa dependió de la "
            "métrica considerada. Por ello, el análisis basado en coeficientes de "
            "variación fue complementado mediante representaciones gráficas y un "
            "índice integrado de variabilidad."
        ),
    )

    for index, (path, description) in enumerate(unique_figures()):
        add_figure(
            doc,
            path,
            description,
            FIGURE_START + index,
            figure_intro(index),
            figure_interpretation(index, tables),
        )

    paragraph(
        doc,
        "La evaluación de la estabilidad permitió identificar configuraciones con "
        "un comportamiento más uniforme durante el periodo de observación. Este "
        "análisis complementó la comparación basada únicamente en valores promedio "
        "y aportó una perspectiva adicional sobre la consistencia operativa del "
        "radioenlace."
    )

    doc.add_heading("4.6. Síntesis de resultados", level=1)
    paragraph(
        doc,
        final_synthesis(
            tables["kruskal"],
            tables["spearman"],
            tables["stability_signal"],
            tables["stability_capacity"],
        ),
    )

    doc.add_page_break()
    doc.add_heading("ANEXO. Comparaciones post hoc de Dunn completas", level=1)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.3)
    sec.right_margin = Cm(1.3)

    add_table(
        doc,
        "dunn_full",
        tables["dunn_full"],
        (
            "La tabla completa permite auditar las comparaciones pareadas y verificar "
            "la coherencia entre las medianas observadas, los estadísticos z, los "
            "valores p ajustados y las decisiones inferenciales."
        ),
    )

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def create_excel(tables: dict[str, pd.DataFrame]) -> None:
    sheets = {
        "Shapiro": tables["shapiro"],
        "Kruskal-Wallis": tables["kruskal"],
        "Dunn resumen": tables["dunn_summary"],
        "Dunn completa": tables["dunn_full"],
        "Descriptivos señal": tables["descriptive_signal"],
        "Descriptivos throughput": tables["descriptive_throughput"],
        "Sensibilidad lluvia": tables["rain_sensitivity"],
        "Conteo lluvia": tables["rain_count"],
        "Lluvia señal": tables["rain_signal"],
        "Lluvia capacidad": tables["rain_capacity"],
        "Spearman horario": tables["spearman"],
        "Estabilidad señal": tables["stability_signal"],
        "Estabilidad capacidad": tables["stability_capacity"],
    }

    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(XLSX_OUT, engine="xlsxwriter") as writer:
        workbook = writer.book
        index = workbook.add_worksheet("Índice")
        header_format = workbook.add_format(
            {"bold": True, "align": "center", "valign": "vcenter", "border": 1}
        )
        index.write_row(0, 0, ["N.º", "Hoja"], header_format)

        for number, (name, df) in enumerate(sheets.items(), start=1):
            sheet = name[:31]
            df_to_write = df.rename(
                columns=lambda c: str(c).replace("Mbps", "Mb/s")
            )
            df_to_write.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes(1, 0)

            for j, col in enumerate(df_to_write.columns):
                series = df_to_write[col].map(
                    lambda x: "" if pd.isna(x) else str(x)
                )
                max_len = int(series.map(len).max()) if not series.empty else 0
                ws.set_column(j, j, min(max(len(str(col)), max_len) + 2, 35))

            index.write(number, 0, number)
            index.write_url(
                number,
                1,
                f"internal:'{sheet}'!A1",
                string=sheet,
            )

        index.set_column("A:A", 7)
        index.set_column("B:B", 28)
